"""Knowledge ingestion and hybrid retrieval.

Published ServiceNow articles and authored runbooks are searchable.
AI drafts stay out of trusted retrieval until a person publishes them.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
from pathlib import Path

from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.config import REPO_ROOT
from app.errors import AppError
from app.models import KnowledgeChunk, KnowledgeDocument, User, utcnow
from app.security import has_permission
from app.util import iso, tokenize

DIM = 256
SEED_DIR = REPO_ROOT / "database" / "seeds" / "knowledge"
TRANSITIONS = {
    "draft": "reviewed",
    "reviewed": "approved",
    "approved": "published",
    "published": "superseded",
}


def embed(text: str) -> list[float]:
    vector = [0.0] * DIM
    for token in tokenize(text):
        digest = hashlib.sha256(token.encode()).digest()
        index = int.from_bytes(digest[:2], "big") % DIM
        sign = 1.0 if digest[2] % 2 == 0 else -1.0
        vector[index] += sign
    norm = math.sqrt(sum(value * value for value in vector)) or 1.0
    return [round(value / norm, 6) for value in vector]


def _cosine(left: list[float], right: list[float]) -> float:
    if not left or not right:
        return 0.0
    return float(sum(a * b for a, b in zip(left, right)))


def chunk_text(text: str, size: int = 700, overlap: int = 80) -> list[str]:
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    if not paragraphs:
        cleaned = text.strip()
        return [cleaned] if cleaned else []
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= size:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= size:
            current = paragraph
            continue
        start = 0
        while start < len(paragraph):
            chunks.append(paragraph[start : start + size])
            start += size - overlap
        current = ""
    if current:
        chunks.append(current)
    return chunks


def extract_upload(filename: str, payload: bytes) -> str:
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension in {"txt", "md", "markdown"}:
        return payload.decode("utf-8", errors="replace")
    if extension == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(payload))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if extension == "docx":
        from docx import Document

        document = Document(io.BytesIO(payload))
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    raise AppError("UNSUPPORTED_DOCUMENT", "Upload a PDF, DOCX, TXT, or Markdown file.", 415)


def upsert_document(
    db: Session,
    *,
    account_id: str,
    title: str,
    content: str,
    source: str,
    doc_type: str,
    status: str,
    application_code: str = "",
    source_id: str = "",
    metadata: dict | None = None,
) -> KnowledgeDocument:
    existing = None
    if source_id:
        existing = (
            db.query(KnowledgeDocument)
            .filter(
                KnowledgeDocument.account_id == account_id,
                KnowledgeDocument.source == source,
                KnowledgeDocument.source_id == source_id,
            )
            .one_or_none()
        )
    if existing is None:
        existing = KnowledgeDocument(
            account_id=account_id,
            title=title,
            content=content,
            source=source,
            doc_type=doc_type,
            status=status,
            application_code=application_code,
            source_id=source_id,
            metadata_json=metadata or {},
        )
        db.add(existing)
        db.flush()
    else:
        existing.title = title
        existing.content = content
        existing.doc_type = doc_type
        existing.status = status
        existing.application_code = application_code
        existing.version = int(existing.version or 1) + 1
        existing.metadata_json = metadata or {}
        existing.updated_at = utcnow()
        db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == existing.id).delete()
    for ordinal, text in enumerate(chunk_text(content)):
        db.add(
            KnowledgeChunk(
                document_id=existing.id,
                ordinal=ordinal,
                text=text,
                embedding=embed(f"{title}\n{text}"),
            )
        )
    db.flush()
    return existing


def load_seed_files(db: Session, account_id: str) -> int:
    if not SEED_DIR.exists():
        return 0
    count = 0
    for path in sorted(SEED_DIR.glob("*.md")):
        meta, body = _frontmatter(path.read_text(encoding="utf-8"))
        upsert_document(
            db,
            account_id=account_id,
            title=meta.get("title") or path.stem,
            content=body,
            source=meta.get("source") or "authored",
            doc_type=meta.get("doc_type") or "article",
            status="published",
            application_code=meta.get("application") or "",
            source_id=meta.get("source_id") or path.stem,
            metadata={"seed_file": path.name, "trusted": True, "authored": True},
        )
        count += 1
    return count


def _frontmatter(text: str) -> tuple[dict, str]:
    if not text.startswith("---"):
        return {}, text.strip()
    parts = text.split("---", 2)
    if len(parts) < 3:
        return {}, text.strip()
    meta = {}
    for line in parts[1].strip().splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip()
    return meta, parts[2].strip()


def search(
    db: Session,
    user: User,
    query: str,
    *,
    application_code: str | None = None,
    limit: int = 5,
    include_drafts: bool = False,
) -> list[dict]:
    if not query.strip():
        return []
    documents = db.query(KnowledgeDocument).filter(KnowledgeDocument.account_id == user.account_id)
    if not include_drafts:
        documents = documents.filter(KnowledgeDocument.status == "published")
    if application_code:
        documents = documents.filter(
            or_(KnowledgeDocument.application_code == application_code, KnowledgeDocument.application_code == "")
        )
    rows = documents.all()
    if not rows:
        return []
    query_tokens = set(tokenize(query))
    query_vector = embed(query)
    ranked: list[dict] = []
    by_id = {row.id: row for row in rows}
    chunks = db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id.in_(list(by_id))).all() if by_id else []
    best: dict[str, dict] = {}
    for chunk in chunks:
        document = by_id.get(chunk.document_id)
        if document is None:
            continue
        text_tokens = set(tokenize(chunk.text + " " + document.title))
        overlap = len(query_tokens & text_tokens) / max(len(query_tokens), 1)
        semantic = _cosine(query_vector, chunk.embedding or [])
        title_bonus = 0.15 if query_tokens & set(tokenize(document.title)) else 0.0
        published_bonus = 0.05 if document.status == "published" else 0.0
        score = (0.55 * overlap) + (0.45 * max(semantic, 0.0)) + title_bonus + published_bonus
        current = best.get(document.id)
        if current is None or score > current["score"]:
            best[document.id] = {
                "score": round(score, 4),
                "document": document,
                "excerpt": chunk.text[:500],
            }
    ranked = sorted(best.values(), key=lambda item: item["score"], reverse=True)
    results = []
    for item in ranked[:limit]:
        if item["score"] <= 0:
            continue
        document = item["document"]
        results.append(
            {
                "id": document.id,
                "title": document.title,
                "source": document.source,
                "source_label": _source_label(document.source),
                "doc_type": document.doc_type,
                "status": document.status,
                "application_code": document.application_code,
                "version": document.version,
                "score": item["score"],
                "excerpt": item["excerpt"],
                "trusted": document.status == "published" and document.source != "ai_draft",
            }
        )
    return results


def _source_label(source: str) -> str:
    return {
        "authored": "Authored runbook",
        "servicenow": "ServiceNow",
        "ai_draft": "AI draft — not approved",
        "upload": "Uploaded document",
    }.get(source, source)


def list_documents(db: Session, user: User, status: str | None = None) -> list[dict]:
    query = db.query(KnowledgeDocument).filter(KnowledgeDocument.account_id == user.account_id)
    if status:
        if status != "published" and not _can_review(user):
            raise AppError("FORBIDDEN", "You can only read published knowledge.", 403)
        query = query.filter(KnowledgeDocument.status == status)
    elif not _can_review(user):
        query = query.filter(KnowledgeDocument.status == "published")
    rows = query.order_by(KnowledgeDocument.updated_at.desc()).all()
    return [_public(row) for row in rows]


def get_document(db: Session, user: User, document_id: str) -> dict:
    row = db.get(KnowledgeDocument, document_id)
    if row is None or row.account_id != user.account_id:
        raise AppError("NOT_FOUND", "Knowledge document was not found.", 404)
    if row.status != "published" and not _can_review(user):
        raise AppError("NOT_FOUND", "Knowledge document was not found.", 404)
    payload = _public(row)
    payload["content"] = row.content
    return payload


def transition(db: Session, user: User, document_id: str, target: str) -> dict:
    row = db.get(KnowledgeDocument, document_id)
    if row is None or row.account_id != user.account_id:
        raise AppError("NOT_FOUND", "Knowledge document was not found.", 404)
    expected = TRANSITIONS.get(row.status)
    if expected != target:
        raise AppError(
            "INVALID_TRANSITION",
            f"A {row.status} article can move to {expected or 'nowhere'}, not {target}.",
            409,
        )
    if target == "reviewed" and not has_permission(user, "knowledge.review"):
        raise AppError("FORBIDDEN", "You cannot review knowledge.", 403)
    if target in {"approved", "published", "superseded"} and not has_permission(user, "knowledge.publish"):
        raise AppError("FORBIDDEN", "You cannot publish knowledge.", 403)
    row.status = target
    row.updated_at = utcnow()
    db.flush()
    return _public(row)


def create_upload(db: Session, user: User, filename: str, payload: bytes, application_code: str = "") -> dict:
    if not has_permission(user, "knowledge.write"):
        raise AppError("FORBIDDEN", "You cannot upload knowledge.", 403)
    text = extract_upload(filename, payload).strip()
    if len(text) < 20:
        raise AppError("EMPTY_DOCUMENT", "No usable text was extracted from that file.", 422)
    title = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ")
    row = upsert_document(
        db,
        account_id=user.account_id,
        title=title,
        content=text,
        source="upload",
        doc_type="upload",
        status="draft",
        application_code=application_code,
        source_id="",
        metadata={"filename": filename, "uploaded_by": user.id},
    )
    return _public(row)


def draft_from_outcome(db: Session, account_id: str, incident, verification, hypothesis) -> KnowledgeDocument | None:
    if not verification or not verification.recovered:
        return None
    existing = (
        db.query(KnowledgeDocument)
        .filter(
            KnowledgeDocument.account_id == account_id,
            KnowledgeDocument.source == "ai_draft",
            KnowledgeDocument.source_id == incident.id,
        )
        .one_or_none()
    )
    statement = hypothesis.statement if hypothesis else "A verified remediation was recorded."
    content = (
        f"# Draft knowledge for {incident.number}\n\n"
        f"This article is an AI draft. It is not approved operational knowledge.\n\n"
        f"Incident: {incident.number}\n\n"
        f"Hypothesis: {statement}\n\n"
        f"Verification: {verification.summary}\n\n"
        "Publish this article only after a person reviews the evidence."
    )
    return upsert_document(
        db,
        account_id=account_id,
        title=f"Draft: verified outcome for {incident.number}",
        content=content,
        source="ai_draft",
        doc_type="ai_draft",
        status="draft",
        application_code="easidoc" if incident.source == "easidoc_simulator" else "",
        source_id=incident.id,
        metadata={"trusted": False, "incident_number": incident.number},
    )


def _can_review(user: User) -> bool:
    return has_permission(user, "knowledge.review") or has_permission(user, "knowledge.publish")


def _public(row: KnowledgeDocument) -> dict:
    return {
        "id": row.id,
        "title": row.title,
        "source": row.source,
        "source_label": _source_label(row.source),
        "doc_type": row.doc_type,
        "status": row.status,
        "application_code": row.application_code,
        "version": row.version,
        "excerpt": (row.content or "")[:280],
        "trusted": row.status == "published" and row.source != "ai_draft",
        "updated_at": iso(row.updated_at),
        "metadata": row.metadata_json or {},
    }
